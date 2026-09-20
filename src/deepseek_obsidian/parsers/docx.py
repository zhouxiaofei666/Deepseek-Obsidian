from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document

from deepseek_obsidian.models import DocumentManifest, PageRecord


class DocxParser:
    name = "python-docx"

    @staticmethod
    def _id(path: Path) -> str:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()[:12]
        stem = re.sub(r"[^A-Za-z0-9._-]+", "-", path.stem).strip("-") or "document"
        return f"{stem}-{digest}"

    def parse(self, source: str | Path, output_root: str | Path) -> DocumentManifest:
        source = Path(source).expanduser().resolve()
        output_root = Path(output_root).expanduser().resolve()
        document_id = self._id(source)
        out = output_root / document_id
        images = out / "images"
        out.mkdir(parents=True, exist_ok=True)
        images.mkdir(exist_ok=True)

        doc = Document(source)
        parts: list[str] = []
        for paragraph in doc.paragraphs:
            value = paragraph.text.strip()
            if value:
                parts.append(value)
        for table_index, table in enumerate(doc.tables, start=1):
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            parts.append(f"\n### Table {table_index}\n")
            parts.append("| " + " | ".join(rows[0]) + " |")
            parts.append("| " + " | ".join("---" for _ in rows[0]) + " |")
            for row in rows[1:]:
                parts.append("| " + " | ".join(row) + " |")

        extracted_images: list[str] = []
        with zipfile.ZipFile(source) as archive:
            for member in archive.namelist():
                if not member.startswith("word/media/") or member.endswith("/"):
                    continue
                name = Path(member).name
                target = images / name
                with archive.open(member) as src, target.open("wb") as dst:
                    shutil.copyfileobj(src, dst)
                extracted_images.append(f"images/{name}")

        text = "\n\n".join(parts)
        for image in extracted_images:
            text += f"\n\n![]({image})"

        page = PageRecord(page=1, text=text, images=extracted_images)
        (out / "document.md").write_text(
            f'---\ntitle: "{source.stem}"\ndocument_id: "{document_id}"\npages: 1\n---\n\n'
            "<!-- PAGE:1 -->\n\n" + text,
            encoding="utf-8",
        )
        (out / "pages.json").write_text(
            json.dumps([page.model_dump()], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        manifest = DocumentManifest(
            document_id=document_id,
            source_file=str(source),
            title=source.stem,
            page_count=1,
            parser=self.name,
            markdown_file="document.md",
            pages_file="pages.json",
            images_dir="images",
            metadata={"source_type": "docx"},
        )
        (out / "document.json").write_text(manifest.model_dump_json(indent=2), encoding="utf-8")
        return manifest
